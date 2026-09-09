# -*- coding: utf-8 -*-
"""xlsx-update-iter003.py - 全景图 xlsx ITER-003 内容回流
- 缺陷清单: 追加 ITER-003 验证快照行（G 列 GitHub 状态同步结论）
- 总览: ITER 状态描述追加 ITER-003
用法: python xlsx-update-iter003.py <xlsx路径>
"""
import sys
from copy import copy
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path)

# --- 缺陷清单 ---
ws = wb['缺陷清单']
# 找表尾
last = ws.max_row
while last > 2 and ws.cell(row=last, column=1).value is None:
    last -= 1
r = last + 1
vals = ['ITER-003', 'P1', '验证', '上游问题验证（2026-09-09）：本批 9 issue #557-565 全部 triage 确认（8 open 转派+#560 修复闭环）；#554 Windows 更新检测 1.1.2-next.4 实测仍未修复（spawnSync npm.cmd 无 shell:true→EINVAL）；#518 README mirror-lag dev 已修（#566）未进 next；#555/#556 无修复待跟踪',
        '跨维度', '验证记录', '未解决 #554（ITER-003 实测复验证据已归档）· #560✅已解决（发布线 608b120 源码实证）· 详见 results/ITER-003-2026-09-09/问题验证记录.md']
src = ws.cell(row=3, column=1)  # 参考样式
for i, v in enumerate(vals, start=1):
    c = ws.cell(row=r, column=i, value=v)
    if i == 1:
        c.font = copy(src.font)
print(f'缺陷清单 追加行 {r}: ITER-003 验证快照')

# --- 总览 ---
ws2 = wb['总览']
updated = 0
for row in ws2.iter_rows():
    for cell in row:
        if cell.value and isinstance(cell.value, str) and '99.2%' in cell.value and '122/123' in cell.value and 'ITER' not in cell.value.replace('ITER-003',''):
            cell.value = cell.value + '（ITER-003：问题验证完成+#554未修复+#560已解决✅+D10-6评测基建建成，基线无漂移）'
            updated += 1
print(f'总览 更新单元格 {updated} 处')

wb.save(path)
print('xlsx 保存完成')